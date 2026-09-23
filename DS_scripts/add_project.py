"""Interactively create OR edit a SISAL project entry: name, status, dates,
people + roles, and the list of SISAL entities it covers.

Edits the CSV source of truth directly (csv/projects.csv, csv/project_person.csv,
csv/project_link_entity.csv) -- NOT the compiled sisalv3.1.db, which is a build
artifact regenerated from these CSVs and would silently discard any direct edit
on the next `build_db.py` run.

Two modes, chosen at startup:
  - New project: name, status, dates, DOI, notes, then add people+roles and
    link entities.
  - Edit existing project: search by id/name, then for each scalar field
    (name/status/dates/DOI/notes) you're asked whether to change it (current
    value shown). People and entities are edited via add/remove/both/skip,
    not just appended to -- entities in particular support removal, since an
    existing project's entity list is exactly the kind of thing that needs
    correcting over time.

When linking entities (new project or "add" during an edit), an ambiguous
name match (multiple candidate entities, can't tell which one from the name
alone) offers a third option beyond picking a number: type 'flag' to send it
to logs/flagged_entity_matches.csv for a data curator to resolve later,
instead of forcing an uncertain guess now. Nothing gets linked for a flagged
entry -- it's deferred, not guessed.

Run from the repo root:
    uv run python DS_scripts/add_project.py
    (or: python3 DS_scripts/add_project.py)

After running, rebuild and verify before committing:
    python3 USER_scripts/build_db.py /tmp/sisal_check
"""
import csv
import re
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

REPO_ROOT = Path(__file__).resolve().parent.parent
CSV_DIR = REPO_ROOT / "csv"

PROJECTS_CSV = CSV_DIR / "projects.csv"
PROJECT_PERSON_CSV = CSV_DIR / "project_person.csv"
PROJECT_LINK_ENTITY_CSV = CSV_DIR / "project_link_entity.csv"
PERSON_CSV = CSV_DIR / "person.csv"
ENTITY_CSV = CSV_DIR / "entity.csv"

FLAG_LOG_CSV = REPO_ROOT / "DS_scripts" / "logs" / "flagged_entity_matches.csv"
FLAG_LOG_FIELDS = [
    "flag_date", "project_id", "project_name", "query",
    "candidate_entity_ids", "candidate_entity_names",
]

PROJECT_STATUS_CHOICES = ["active", "closed"]
# Same contributor_role_enum used by release_person -- kept in sync manually,
# see ROLE in build_db.py.
ROLE_CHOICES = [
    "release_steward",
    "data_curator",
    "workflow_developer",
    "project_lead",
    "data_contributor",
]

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


# ---------------------------------------------------------------- CSV I/O

def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return reader.fieldnames, list(reader)


def write_csv(path, fieldnames, rows, use_bom=False):
    encoding = "utf-8-sig" if use_bom else "utf-8"
    with open(path, "w", newline="", encoding=encoding) as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def log_flagged_match(project_context, query, matches):
    """Append one ambiguous entity match to the flagged-for-review log
    instead of guessing. Creates the log (with header) if it doesn't exist."""
    FLAG_LOG_CSV.parent.mkdir(parents=True, exist_ok=True)
    is_new = not FLAG_LOG_CSV.exists()
    with open(FLAG_LOG_CSV, "a", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=FLAG_LOG_FIELDS)
        if is_new:
            writer.writeheader()
        writer.writerow({
            "flag_date": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "project_id": project_context.get("project_id", ""),
            "project_name": project_context.get("project_name", ""),
            "query": query,
            "candidate_entity_ids": ";".join(m["entity_id"] for m in matches),
            "candidate_entity_names": ";".join(m["entity_name"] for m in matches),
        })


# ---------------------------------------------------------------- prompts

def prompt(text, required=True, default=None):
    suffix = f" [{default}]" if default is not None else ""
    while True:
        val = input(f"{text}{suffix}: ").strip()
        if not val and default is not None:
            return default
        if not val and not required:
            return ""
        if val:
            return val
        print("  This is required -- please enter a value.")


def prompt_choice(text, choices, default=None):
    choice_str = "/".join(choices)
    while True:
        val = prompt(f"{text} ({choice_str})", required=default is None, default=default)
        val = val.lower()
        if val in choices:
            return val
        print(f"  Please enter one of: {choice_str}")


def prompt_date(text, required=True):
    while True:
        val = input(f"{text} (YYYY-MM-DD, blank to skip): ").strip()
        if not val:
            if required:
                print("  This is required -- please enter a date.")
                continue
            return ""
        if not DATE_RE.match(val):
            print("  Please use YYYY-MM-DD format.")
            continue
        try:
            datetime.strptime(val, "%Y-%m-%d")
        except ValueError:
            print("  Not a real calendar date -- try again.")
            continue
        return val


def prompt_yes_no(text, default_yes=True):
    default = "y" if default_yes else "n"
    while True:
        val = input(f"{text} (y/n) [{default}]: ").strip().lower()
        if not val:
            val = default
        if val in ("y", "yes"):
            return True
        if val in ("n", "no"):
            return False
        print("  Please answer y or n.")


# ---------------------------------------------------------------- people

def find_person(people_rows, query):
    """Case-insensitive substring match on name. Returns list of matching rows."""
    q = query.strip().lower()
    return [r for r in people_rows if q in r["name"].lower()]


def add_people_loop(people_rows, existing_pairs=None):
    """Interactively add people+roles. `existing_pairs` (a set of
    (person_id, role) tuples already on the project) prevents adding an exact
    duplicate -- project_person's primary key is (project_id, person_id, role),
    so a literal duplicate would fail at build time."""
    existing_pairs = set(existing_pairs) if existing_pairs else set()
    print("\n--- People to add ---")
    print("For each person, search by (partial) name; pick from matches or add a new person.")
    collected = []  # list of (person_id, name, role)
    while True:
        query = prompt("\nSearch for a person by name (blank to finish)", required=False)
        if not query:
            break
        matches = find_person(people_rows, query)
        if not matches:
            print(f"  No existing person matches '{query}'.")
            if prompt_yes_no(f"  Add '{query}' as a brand-new person?", default_yes=True):
                new_id = str(max((int(r["person_id"]) for r in people_rows), default=0) + 1)
                orcid = prompt("  ORCID (blank if unknown)", required=False)
                new_row = {"person_id": new_id, "name": query, "orcid": orcid}
                people_rows.append(new_row)
                chosen = new_row
            else:
                continue
        elif len(matches) == 1:
            chosen = matches[0]
            print(f"  Matched: {chosen['name']} (person_id {chosen['person_id']}, ORCID {chosen['orcid'] or 'unknown'})")
        else:
            print(f"  {len(matches)} matches:")
            for i, m in enumerate(matches, 1):
                print(f"    {i}. {m['name']} (person_id {m['person_id']}, ORCID {m['orcid'] or 'unknown'})")
            while True:
                idx = prompt("  Pick a number")
                if idx.isdigit() and 1 <= int(idx) <= len(matches):
                    chosen = matches[int(idx) - 1]
                    break
                print("  Invalid choice.")

        role = prompt_choice("  Role for this person on this project", ROLE_CHOICES)
        if (chosen["person_id"], role) in existing_pairs:
            print(f"  {chosen['name']} is already on this project as {role} -- skipping.")
            continue
        existing_pairs.add((chosen["person_id"], role))
        collected.append((chosen["person_id"], chosen["name"], role))
        print(f"  Added: {chosen['name']} as {role}")

    return collected


def remove_people_flow(current_people):
    """current_people: list of dicts with person_id, name, role. Returns the
    subset the user picked to remove."""
    if not current_people:
        print("  No people currently on this project.")
        return []
    print("\nCurrent people on this project:")
    for i, p in enumerate(current_people, 1):
        print(f"  {i}. {p['name']} (person_id {p['person_id']}) -- {p['role']}")
    raw = prompt("Numbers to remove (comma-separated, blank to remove none)", required=False)
    if not raw:
        return []
    to_remove = []
    for tok in raw.split(","):
        tok = tok.strip()
        if tok.isdigit() and 1 <= int(tok) <= len(current_people):
            to_remove.append(current_people[int(tok) - 1])
        elif tok:
            print(f"  Ignoring invalid selection '{tok}'.")
    for p in to_remove:
        print(f"  Will remove: {p['name']} -- {p['role']}")
    return to_remove


# ---------------------------------------------------------------- entities

def find_entities(entity_rows, query):
    """Match entity_id exactly; else entity_name exactly (case-insensitive); else
    entity_name by substring. Each tier only falls through to the next if it finds
    nothing -- otherwise a short, valid, exact name (e.g. "A1") would get swallowed
    into unrelated substring collisions ("SPA121" also contains "a1") instead of
    matching the entity that's actually named exactly that."""
    q = query.strip().lower()
    exact_id = [r for r in entity_rows if r["entity_id"] == query.strip()]
    if exact_id:
        return exact_id
    exact_name = [r for r in entity_rows if (r["entity_name"] or "").strip().lower() == q]
    if exact_name:
        return exact_name
    return [r for r in entity_rows if q in (r["entity_name"] or "").lower()]


def read_tokens_from_file(path_str):
    """Read a list of entity id/name tokens from column A of an .xlsx or .csv/.txt
    file. Returns a list of non-empty string tokens, or None if the file couldn't
    be read (caller should fall back to another input method)."""
    path = Path(path_str.strip().strip('"').strip("'")).expanduser()
    if not path.exists():
        print(f"  File not found: {path}")
        return None

    raw_values = []
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        if openpyxl is None:
            print("  openpyxl isn't installed -- can't read .xlsx. Try a .csv/.txt export instead.")
            return None
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        if len(wb.sheetnames) > 1:
            print(f"  Note: workbook has {len(wb.sheetnames)} sheets, reading the first one ('{ws.title}').")
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            val = row[0]
            if val is not None and str(val).strip():
                raw_values.append(str(val).strip())
        wb.close()
    elif suffix in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.reader(f, delimiter=delimiter):
                if row and row[0].strip():
                    raw_values.append(row[0].strip())
    else:
        print(f"  Unrecognized file type '{suffix}' -- expected .xlsx, .csv, .tsv, or .txt.")
        return None

    if not raw_values:
        print("  No values found in column A of that file.")
        return None

    print(f"\n  Read {len(raw_values)} value(s) from column A, first 5:")
    for v in raw_values[:5]:
        print(f"    - {v}")
    if prompt_yes_no("  Does the first row look like a header (e.g. 'entity_name') rather than real data?", default_yes=False):
        raw_values = raw_values[1:]
        print(f"  Skipping header, {len(raw_values)} value(s) remain.")

    return raw_values


def resolve_entity_token(entity_rows, query, seen_ids, linked, project_context):
    """Resolve one entity_id/name token interactively (disambiguating if needed)
    and append it to `linked` in place if found and not already added. On an
    ambiguous match, offers 'flag' as an alternative to picking -- logs the
    query and its candidates for a data curator to resolve later, and leaves
    it unlinked rather than guessing."""
    matches = find_entities(entity_rows, query)
    if not matches:
        print(f"  No entity matches '{query}' -- check the ID/name and try again.")
        return
    if len(matches) == 1:
        chosen = matches[0]
    else:
        print(f"  '{query}' -- {len(matches)} matches:")
        for i, m in enumerate(matches, 1):
            print(f"    {i}. entity_id {m['entity_id']} -- {m['entity_name']}")
        while True:
            idx = prompt("  Pick a number, or type 'flag' to send this to the data curator instead")
            if idx.strip().lower() in ("flag", "f"):
                log_flagged_match(project_context, query, matches)
                print(f"  Flagged '{query}' for the data curator (see {FLAG_LOG_CSV.name}) -- not linked for now.")
                return
            if idx.isdigit() and 1 <= int(idx) <= len(matches):
                chosen = matches[int(idx) - 1]
                break
            print("  Invalid choice.")

    if chosen["entity_id"] in seen_ids:
        print(f"  entity_id {chosen['entity_id']} already linked, skipping.")
        return
    seen_ids.add(chosen["entity_id"])
    linked.append((chosen["entity_id"], chosen["entity_name"]))
    print(f"  Added: entity_id {chosen['entity_id']} ({chosen['entity_name']})")


def collect_entities(entity_rows, project_context, seen_ids=None):
    """Interactively link entities: from a file, pasted list, and/or one at a
    time. `seen_ids` (if given) is treated as already-linked -- used during an
    edit so re-adding an existing link is reported and skipped, not duplicated."""
    print("\n--- SISAL entities to link ---")
    linked = []  # list of (entity_id, entity_name)
    seen_ids = set(seen_ids) if seen_ids else set()

    if prompt_yes_no("Do you have a file with the list (Excel .xlsx or .csv/.txt)?", default_yes=True):
        while True:
            file_path = prompt("  Path to the file", required=False)
            if not file_path:
                break
            tokens = read_tokens_from_file(file_path)
            if tokens is not None:
                print(f"\nResolving {len(tokens)} entries from the file...")
                for tok in tokens:
                    resolve_entity_token(entity_rows, tok, seen_ids, linked, project_context)
                break
            if not prompt_yes_no("  Try a different file path?", default_yes=True):
                break

    print("\nYou can also paste a list directly (comma- and/or newline-separated,")
    print("entity IDs or entity names, mixed is fine) -- press Enter to skip this.")
    print("(End a multi-line paste with an empty line.)")
    pasted_lines = []
    while True:
        line = input("> ")
        if not line.strip():
            break
        pasted_lines.append(line)

    if pasted_lines:
        tokens = []
        for line in pasted_lines:
            tokens.extend(t.strip() for t in line.split(",") if t.strip())
        print(f"\nResolving {len(tokens)} pasted entries...")
        for tok in tokens:
            resolve_entity_token(entity_rows, tok, seen_ids, linked, project_context)

    print("\nAdd more one at a time (blank line to finish), or just press Enter if you're done.")
    while True:
        query = prompt("\nEntity (id or name, blank to finish)", required=False)
        if not query:
            break
        resolve_entity_token(entity_rows, query, seen_ids, linked, project_context)

    return linked


def resolve_removal_token(current_linked, query, seen_ids, to_remove):
    """Same matching tiers as resolve_entity_token, but scoped to the project's
    *currently linked* entities only -- so removal can't accidentally match some
    unrelated entity elsewhere in the database. No flag option here; removal
    ambiguity is bounded to this project's own (much smaller) entity list."""
    matches = find_entities(current_linked, query)
    if not matches:
        print(f"  '{query}' is not currently linked to this project.")
        return
    if len(matches) == 1:
        chosen = matches[0]
    else:
        print(f"  '{query}' -- {len(matches)} matches among linked entities:")
        for i, m in enumerate(matches, 1):
            print(f"    {i}. entity_id {m['entity_id']} -- {m['entity_name']}")
        while True:
            idx = prompt("  Pick a number")
            if idx.isdigit() and 1 <= int(idx) <= len(matches):
                chosen = matches[int(idx) - 1]
                break
            print("  Invalid choice.")

    if chosen["entity_id"] in seen_ids:
        return
    seen_ids.add(chosen["entity_id"])
    to_remove.append(chosen)
    print(f"  Will remove: entity_id {chosen['entity_id']} ({chosen['entity_name']})")


def collect_entities_to_remove(current_linked):
    """current_linked: list of dicts with entity_id, entity_name (the project's
    current links). Returns the subset picked for removal."""
    print(f"\nThis project currently has {len(current_linked)} linked entities.")
    if current_linked and prompt_yes_no("Show the full list before choosing what to remove?", default_yes=False):
        for e in current_linked:
            print(f"  entity_id {e['entity_id']} -- {e['entity_name']}")

    print("\nEnter entity IDs/names to remove -- paste a list (comma/newline separated)")
    print("and/or add one at a time. (Blank line to finish each step.)")
    to_remove = []
    seen_ids = set()

    pasted_lines = []
    while True:
        line = input("> ")
        if not line.strip():
            break
        pasted_lines.append(line)
    tokens = []
    for line in pasted_lines:
        tokens.extend(t.strip() for t in line.split(",") if t.strip())
    for tok in tokens:
        resolve_removal_token(current_linked, tok, seen_ids, to_remove)

    while True:
        query = prompt("\nEntity to remove (id or name, blank to finish)", required=False)
        if not query:
            break
        resolve_removal_token(current_linked, query, seen_ids, to_remove)

    return to_remove


# ---------------------------------------------------------------- projects

def find_and_pick_project(proj_rows):
    if not proj_rows:
        print("No projects exist yet -- nothing to edit.")
        return None
    while True:
        query = prompt("\nSearch for the project (id or partial name, blank to cancel)", required=False)
        if not query:
            return None
        q = query.strip().lower()
        exact_id = [r for r in proj_rows if r["project_id"] == query.strip()]
        matches = exact_id or [r for r in proj_rows if q in (r["project_name"] or "").lower()]
        if not matches:
            print(f"  No project matches '{query}'.")
            continue
        if len(matches) == 1:
            return matches[0]
        print(f"  {len(matches)} matches:")
        for i, m in enumerate(matches, 1):
            print(f"    {i}. project_id {m['project_id']} -- {m['project_name']} ({m['project_status']})")
        while True:
            idx = prompt("  Pick a number")
            if idx.isdigit() and 1 <= int(idx) <= len(matches):
                return matches[int(idx) - 1]
            print("  Invalid choice.")


def create_project(proj_fields, proj_rows, pp_fields, pp_rows, ple_fields, ple_rows,
                    person_fields, person_rows, entity_fields, entity_rows):
    print("\n=== New SISAL project entry ===\n")

    project_name = prompt("Project name")
    new_project_id = str(max((int(r["project_id"]) for r in proj_rows), default=0) + 1)
    project_context = {"project_id": new_project_id, "project_name": project_name}

    status = prompt_choice("Project status", PROJECT_STATUS_CHOICES, default="active")
    start_date = prompt_date("Start date")
    if status == "closed":
        completion_date = prompt_date("Completion date")
    else:
        if prompt_yes_no("Add a (planned/target) completion date anyway?", default_yes=False):
            completion_date = prompt_date("Completion date", required=False)
        else:
            completion_date = ""

    project_doi = prompt("Project DOI (blank if none yet)", required=False)
    project_notes = prompt("Project notes (blank to skip)", required=False)

    people = add_people_loop(person_rows)
    entities = collect_entities(entity_rows, project_context)

    print("\n=== Summary ===")
    print(f"project_id:       {new_project_id}")
    print(f"project_name:     {project_name}")
    print(f"project_status:   {status}")
    print(f"start_date:       {start_date or '(none)'}")
    print(f"completion_date:  {completion_date or '(none)'}")
    print(f"project_doi:      {project_doi or '(none)'}")
    print(f"project_notes:    {project_notes or '(none)'}")
    print(f"people ({len(people)}):")
    for pid, name, role in people:
        print(f"  - {name} (person_id {pid}) -- {role}")
    print(f"linked entities ({len(entities)}):")
    for eid, ename in entities:
        print(f"  - entity_id {eid} ({ename})")

    if not prompt_yes_no("\nWrite this to the CSVs?", default_yes=True):
        print("Aborted -- nothing written.")
        return

    proj_rows.append({
        "project_id": new_project_id,
        "project_name": project_name,
        "project_doi": project_doi,
        "project_status": status,
        "start_date": start_date,
        "completion_date": completion_date,
        "project_notes": project_notes,
    })
    for pid, _name, role in people:
        pp_rows.append({"project_id": new_project_id, "person_id": pid, "role": role})
    for eid, _ename in entities:
        ple_rows.append({"project_id": new_project_id, "entity_id": eid})

    write_csv(PROJECTS_CSV, proj_fields, proj_rows, use_bom=True)
    write_csv(PROJECT_PERSON_CSV, pp_fields, pp_rows)
    write_csv(PROJECT_LINK_ENTITY_CSV, ple_fields, ple_rows)
    # person.csv may have gained a brand-new person row -- keep its BOM (Excel-safe accents).
    write_csv(PERSON_CSV, person_fields, person_rows, use_bom=True)

    print(f"\nDone. Project '{project_name}' added as project_id {new_project_id}.")
    print("Next: rebuild and verify before committing:")
    print("  python3 USER_scripts/build_db.py /tmp/sisal_check")


def edit_project(project, proj_fields, proj_rows, pp_fields, pp_rows, ple_fields, ple_rows,
                  person_fields, person_rows, entity_fields, entity_rows):
    project_id = project["project_id"]
    project_context = {"project_id": project_id, "project_name": project["project_name"]}

    print(f"\n=== Editing project_id {project_id}: {project['project_name']} ===")
    print(f"  status:          {project['project_status']}")
    print(f"  start_date:      {project['start_date'] or '(none)'}")
    print(f"  completion_date: {project['completion_date'] or '(none)'}")
    print(f"  project_doi:     {project['project_doi'] or '(none)'}")
    print(f"  project_notes:   {project['project_notes'] or '(none)'}")

    current_people = [
        {
            "person_id": pp["person_id"],
            "role": pp["role"],
            "name": next((p["name"] for p in person_rows if p["person_id"] == pp["person_id"]), "?"),
        }
        for pp in pp_rows if pp["project_id"] == project_id
    ]
    print(f"  people ({len(current_people)}):")
    for p in current_people:
        print(f"    - {p['name']} (person_id {p['person_id']}) -- {p['role']}")

    current_linked = [
        {
            "entity_id": e["entity_id"],
            "entity_name": next((x["entity_name"] for x in entity_rows if x["entity_id"] == e["entity_id"]), "?"),
        }
        for e in ple_rows if e["project_id"] == project_id
    ]
    print(f"  linked entities: {len(current_linked)}")

    # --- scalar fields: ask per field whether to change it ---
    changes = {}
    if prompt_yes_no(f"\nChange project_name? current: '{project['project_name']}'", default_yes=False):
        changes["project_name"] = prompt("  New project_name", required=False)
    if prompt_yes_no(f"Change project_status? current: '{project['project_status']}'", default_yes=False):
        changes["project_status"] = prompt_choice("  New status", PROJECT_STATUS_CHOICES)
    if prompt_yes_no(f"Change start_date? current: '{project['start_date'] or '(none)'}'", default_yes=False):
        changes["start_date"] = prompt_date("  New start_date", required=False)
    if prompt_yes_no(f"Change completion_date? current: '{project['completion_date'] or '(none)'}'", default_yes=False):
        changes["completion_date"] = prompt_date("  New completion_date", required=False)
    if prompt_yes_no(f"Change project_doi? current: '{project['project_doi'] or '(none)'}'", default_yes=False):
        changes["project_doi"] = prompt("  New project_doi", required=False)
    if prompt_yes_no(f"Change project_notes? current: '{project['project_notes'] or '(none)'}'", default_yes=False):
        changes["project_notes"] = prompt("  New project_notes", required=False)

    # --- people: add / remove / both / skip ---
    people_to_add, people_to_remove = [], []
    people_action = prompt_choice("\nPeople -- add, remove, both, or skip?", ["add", "remove", "both", "skip"], default="skip")
    if people_action in ("add", "both"):
        existing_pairs = {(p["person_id"], p["role"]) for p in current_people}
        people_to_add = add_people_loop(person_rows, existing_pairs)
    if people_action in ("remove", "both"):
        people_to_remove = remove_people_flow(current_people)

    # --- entities: add / remove / both / skip ---
    entities_to_add, entities_to_remove = [], []
    entity_action = prompt_choice("\nEntities -- add, remove, both, or skip?", ["add", "remove", "both", "skip"], default="skip")
    if entity_action in ("add", "both"):
        seen_ids = {e["entity_id"] for e in current_linked}
        entities_to_add = collect_entities(entity_rows, project_context, seen_ids=seen_ids)
    if entity_action in ("remove", "both"):
        entities_to_remove = collect_entities_to_remove(current_linked)

    # --- summary ---
    print("\n=== Summary of changes ===")
    if changes:
        for k, v in changes.items():
            print(f"  {k}: '{project[k]}' -> '{v}'")
    else:
        print("  (no field changes)")
    print(f"  people to add ({len(people_to_add)}):")
    for pid, name, role in people_to_add:
        print(f"    + {name} -- {role}")
    print(f"  people to remove ({len(people_to_remove)}):")
    for p in people_to_remove:
        print(f"    - {p['name']} -- {p['role']}")
    print(f"  entities to add ({len(entities_to_add)}):")
    for eid, ename in entities_to_add[:10]:
        print(f"    + entity_id {eid} ({ename})")
    if len(entities_to_add) > 10:
        print(f"    ... and {len(entities_to_add) - 10} more")
    print(f"  entities to remove ({len(entities_to_remove)}):")
    for e in entities_to_remove[:10]:
        print(f"    - entity_id {e['entity_id']} ({e['entity_name']})")
    if len(entities_to_remove) > 10:
        print(f"    ... and {len(entities_to_remove) - 10} more")

    if not (changes or people_to_add or people_to_remove or entities_to_add or entities_to_remove):
        print("\nNo changes made -- nothing to write.")
        return

    if not prompt_yes_no("\nWrite these changes to the CSVs?", default_yes=True):
        print("Aborted -- nothing written.")
        return

    # --- apply ---
    for k, v in changes.items():
        project[k] = v  # `project` is the live dict inside proj_rows -- mutating it in place is enough

    for pid, _name, role in people_to_add:
        pp_rows.append({"project_id": project_id, "person_id": pid, "role": role})
    remove_pairs = {(p["person_id"], p["role"]) for p in people_to_remove}
    pp_rows[:] = [r for r in pp_rows if not (r["project_id"] == project_id and (r["person_id"], r["role"]) in remove_pairs)]

    for eid, _ename in entities_to_add:
        ple_rows.append({"project_id": project_id, "entity_id": eid})
    remove_eids = {e["entity_id"] for e in entities_to_remove}
    ple_rows[:] = [r for r in ple_rows if not (r["project_id"] == project_id and r["entity_id"] in remove_eids)]

    write_csv(PROJECTS_CSV, proj_fields, proj_rows, use_bom=True)
    write_csv(PROJECT_PERSON_CSV, pp_fields, pp_rows)
    write_csv(PROJECT_LINK_ENTITY_CSV, ple_fields, ple_rows)
    write_csv(PERSON_CSV, person_fields, person_rows, use_bom=True)

    print(f"\nDone. project_id {project_id} updated.")
    print("Next: rebuild and verify before committing:")
    print("  python3 USER_scripts/build_db.py /tmp/sisal_check")


def main():
    proj_fields, proj_rows = read_csv(PROJECTS_CSV)
    pp_fields, pp_rows = read_csv(PROJECT_PERSON_CSV)
    ple_fields, ple_rows = read_csv(PROJECT_LINK_ENTITY_CSV)
    person_fields, person_rows = read_csv(PERSON_CSV)
    entity_fields, entity_rows = read_csv(ENTITY_CSV)

    print("=== SISAL project entry ===")
    mode = prompt_choice("\nCreate a new project, or edit an existing one?", ["new", "edit"], default="new")

    if mode == "edit":
        project = find_and_pick_project(proj_rows)
        if project is None:
            print("No project selected -- nothing to do.")
            return
        edit_project(
            project, proj_fields, proj_rows, pp_fields, pp_rows, ple_fields, ple_rows,
            person_fields, person_rows, entity_fields, entity_rows,
        )
    else:
        create_project(
            proj_fields, proj_rows, pp_fields, pp_rows, ple_fields, ple_rows,
            person_fields, person_rows, entity_fields, entity_rows,
        )


if __name__ == "__main__":
    main()
